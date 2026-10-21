#!/usr/bin/env python3
"""Build the review copy from docs/DEVELOPMENT-PLAN.md using python-docx.

Run with the document runtime Python. Render the output with the documents
skill's render_docx.py and review every page before committing the DOCX.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT = "Arial Unicode MS"
INK = "000000"
BORDER = "D9D9D9"
ROOT = Path(__file__).resolve().parents[2]
INLINE = re.compile(r"(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def font_properties(run, size=None):
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(INK)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{name}"), None)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)
    if size is not None:
        run.font.size = Pt(size)


def inline(paragraph, value):
    for token in INLINE.split(value):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
        elif token.startswith("[") and "](" in token:
            label, target = token[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            # Relative source references remain readable in a detached DOCX.
            if target.endswith(".md") and label != target.removesuffix(".md"):
                paragraph.add_run(f"（{target}）")
        else:
            run = paragraph.add_run(token)
        font_properties(run)


def clean_heading(value):
    return re.sub(r"^\d+\.\d+\s+", "", value)


def bookmark(paragraph, number):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), f"chapter_{number}")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def navigation(doc, headings):
    doc.add_paragraph("阅读导航", "Heading 2")
    for number, title in enumerate(headings, 1):
        p = doc.add_paragraph(style="Contents Entry")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), f"chapter_{number}")
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for key in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn(f"w:{key}"), FONT)
        props.append(fonts)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "21")
        props.append(size)
        run.append(props)
        text = OxmlElement("w:t")
        text.text = title
        run.append(text)
        link.append(run)
        p._p.append(link)
    doc.add_page_break()


def table(doc, lines):
    rows = [[x.strip() for x in row.strip().strip("|").split("|")] for row in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", x) for x in row)]
    count = len(rows[0])
    fractions = {
        3: (0.19, 0.49, 0.32),
        4: (0.16, 0.28, 0.36, 0.20),
        5: (0.14, 0.38, 0.18, 0.18, 0.12),
    }.get(count, tuple(1 / count for _ in range(count)))
    t = doc.add_table(rows=0, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", BORDER)):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    props.append(borders)
    for col, fraction in zip(t.columns, fractions):
        col.width = Inches(7.0 * fraction)
    for index, values in enumerate(rows):
        row = t.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            header = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(header)
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = Inches(7.0 * fractions[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "173650" if index == 0 else ("F3F6F9" if index % 2 == 0 else "FFFFFF"))
            cp.append(shade)
            margins = OxmlElement("w:tcMar")
            for name, amount in (("top", "90"), ("bottom", "90"), ("left", "100"), ("right", "100")):
                side = OxmlElement(f"w:{name}")
                side.set(qn("w:w"), amount)
                side.set(qn("w:type"), "dxa")
                margins.append(side)
            cp.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            if count == 5 and column == 4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value)
            for run in p.runs:
                font_properties(run, 10.5)
                if index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.35
    spacer.add_run().font.size = Pt(3)


def build(source, output):
    text = source.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()
    # Bundled Word templates can carry a title rule; remove inherited borders.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.70)
    section.bottom_margin = Inches(0.68)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.30)
    section.different_first_page_header_footer = True
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.color.rgb = RGBColor.from_string(INK)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{key}"), None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (("Title", 27, 0, 16), ("Heading 1", 17, 17, 9), ("Heading 2", 12.5, 12, 6)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    contents = doc.styles.add_style("Contents Entry", 1)
    contents.base_style = normal
    contents.font.size = Pt(10.5)
    contents.paragraph_format.space_after = Pt(2)
    contents.paragraph_format.line_spacing = 1.10
    contents.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.style.font.size = Pt(9)
    doc.core_properties.title = "UniPO 开发规划书"
    doc.core_properties.subject = "单校试点开发范围与实施计划"
    doc.core_properties.author = "UniPO"
    doc.core_properties.keywords = "UniPO,开发规划,产品,设计,工程,运营"
    doc.core_properties.comments = ""
    headings = [line[3:].strip() for line in lines if line.startswith("## ")]
    chapter = 0
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("| "):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            table(doc, block)
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level = len(match[1])
            if level == 2 and chapter == 0:
                navigation(doc, headings)
            p = doc.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            inline(p, clean_heading(match[2]))
            if level == 2:
                chapter += 1
                bookmark(p, chapter)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            inline(p, line)
        index += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output} with {chapter} chapters and {len(doc.tables)} tables")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/DEVELOPMENT-PLAN.md")
    parser.add_argument("--output", type=Path, default=ROOT / "docs/UniPO-开发规划书-v1.0.docx")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
